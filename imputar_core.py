"""
Motor de imputación. Importable desde app.py (Streamlit) o scripts CLI.
No tiene side effects al importar.
"""

import io
import re
import datetime
import openpyxl
from openpyxl.styles import PatternFill

from mep_helper import mep_para_fecha

YELLOW_FILL = PatternFill(patternType='solid', fgColor='FFFF00')

MESES_ES = {
    1: 'enero', 2: 'febrero', 3: 'marzo', 4: 'abril',
    5: 'mayo',  6: 'junio',   7: 'julio', 8: 'agosto',
    9: 'septiembre', 10: 'octubre', 11: 'noviembre', 12: 'diciembre',
}

SHEETS_BASE_PESOS = {
    'INDICE CAC': {
        'header_row': 5, 'data_start': 6,
        'cuit_col': 12, 'nombre_col': 9, 'lote_col': 8,
    },
    'INDICE CAC M. OBRA': {
        'header_row': 3, 'data_start': 4,
        'cuit_col': 11, 'nombre_col': 9, 'lote_col': 8,
    },
    'BOLSA CEMENTO': {
        'header_row': 3, 'data_start': 4,
        'cuit_col': 14, 'nombre_col': 10, 'lote_col': 9,
    },
}

SHEETS_BASE_USD = {
    '$  USD fijo': {
        'header_row': 3, 'data_start': 4,
        'cuit_col': 13, 'nombre_col': 10, 'lote_col': 9,
    },
}


def normalize_cuits(raw):
    """Extrae todos los CUITs de una celda de deudores.

    Maneja guiones/puntos (20-11111111-2), varios CUITs separados por '/'
    y celdas con texto libre entre CUITs
    ('27-11111111-2 (nombre) apellido 20-33333333-4' → dos CUITs).
    """
    if raw is None or raw == '':
        return []
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    s = str(raw)
    # Unir grupos de dígitos conectados por guion/punto (con espacios sueltos
    # alrededor) mientras no excedan los 12 dígitos de un CUIT. Un guion entre
    # dos CUITs completos ("20111111112 - 27333333334") no los une porque el
    # resultado se pasaría de largo.
    results, cur, prev_end = [], '', None
    for g in re.finditer(r'\d+', s):
        sep = s[prev_end:g.start()] if prev_end is not None else None
        unido = sep is not None and re.fullmatch(r'\s*[.\-]\s*', sep)
        if cur and unido and len(cur) + len(g.group()) <= 12:
            cur += g.group()
        else:
            if 10 <= len(cur) <= 12:
                results.append(cur)
            cur = g.group()
        prev_end = g.end()
    if 10 <= len(cur) <= 12:
        results.append(cur)
    if not results:
        # formato raro (ej: dígitos separados solo por espacios): unir todo
        for part in s.split('/'):
            digits = re.sub(r'\D', '', part)
            if len(digits) >= 10:
                results.append(digits)
    return list(dict.fromkeys(results))


def is_row_yellow(ws, row_num):
    for cell in ws[row_num]:
        if cell.value is None:
            continue
        fill = cell.fill
        if fill and fill.patternType and fill.patternType != 'none':
            rgb = str(fill.fgColor.rgb) if fill.fgColor else ''
            if 'FFFF00' in rgb or rgb == 'FFFFFF00':
                return True
        break
    return False


def extract_cuit_from_concepto(concepto):
    if not concepto:
        return None
    matches = re.findall(r'\b(\d{11,12})\b', str(concepto))
    return matches[0] if matches else None


def parse_date(val):
    if isinstance(val, datetime.datetime):
        return val
    if isinstance(val, str):
        for fmt in ('%d-%m-%Y', '%Y-%m-%d', '%d/%m/%Y'):
            try:
                return datetime.datetime.strptime(val.strip(), fmt)
            except ValueError:
                continue
    return None


def _norm(s):
    return s.lower().replace('ó', 'o').replace('é', 'e').replace('á', 'a').replace('í', 'i').replace('ú', 'u')


def max_cuota_celda(val):
    """Máxima cuota implicada por una celda de "NUMERO DE CUOTA".

    Acepta números, listas tipo "10 y 11" / "2 Y 3" / "3, 4 y 5" (el formato que
    escribe el propio bot al imputar varias cuotas juntas) y "parte de cX"
    (la cuota X está parcialmente paga → el siguiente pago completo es X+1).
    Otros textos ("1 al 12 inclusive", etc.) se ignoran por ambiguos.
    """
    if isinstance(val, (int, float)):
        n = int(val)
        return n if 0 < n <= 200 else None
    if isinstance(val, str):
        s = val.strip()
        if 'parte' in s.lower():
            # "17 y parte de c18" → 18 (la 18 quedó a medias: el próximo pago
            # entero es la 19). Las notas entre paréntesis no cuentan.
            nums = [int(x) for x in re.findall(r'\d+', _sin_notas(s))]
        elif re.fullmatch(r'[cC]?\s*\d+(\s*[,yY]\s*[cC]?\s*\d+)*\.?', s):
            nums = [int(x) for x in re.findall(r'\d+', s)]
        else:
            return None
        nums = [n for n in nums if 0 < n <= 200]
        return max(nums) if nums else None
    return None


def ultima_fila_con_datos(ws, cols=None):
    """Última fila de `ws` que realmente tiene un valor (en `cols` si se pasan).

    `ws.max_row` puede dar 1.048.576 en hojas con formato fantasma (celdas
    vacías pero pintadas/con estilo). Recorrer hasta ahí con `ws.cell()` o
    `iter_rows()` MATERIALIZA millones de celdas: una corrida USD llegaba a
    ~7 GB de RAM y Streamlit Cloud mataba el proceso. Acá miramos las celdas
    que openpyxl ya tiene cargadas, sin crear ninguna nueva.
    """
    cells = getattr(ws, '_cells', None)
    if not cells:
        return ws.max_row
    if cols:
        cols = set(cols)
        filas = [r for (r, c), cell in cells.items()
                 if c in cols and cell.value is not None]
    else:
        filas = [r for (r, c), cell in cells.items() if cell.value is not None]
    return max(filas) if filas else 0


def detectar_mes_transferencias(ws_imp, max_row=500):
    """Lee col A de imputaciones y retorna (year, month) más frecuente, o (None, None).

    Solo cuenta filas PENDIENTES (no amarillas): las hojas —sobre todo las USD—
    son acumulativas y arrastran transferencias ya imputadas de meses anteriores.
    Contar todas hacía que el mes detectado fuera un mes viejo ya imputado y el
    script reportara "mes ya imputado" para todo sin imputar nada.
    Si no hubiera filas pendientes, cae a contar todas (comportamiento previo).
    """
    from collections import Counter
    conteo = Counter()
    conteo_todas = Counter()
    for row in ws_imp.iter_rows(min_row=4, max_row=max_row, max_col=1):
        dt = parse_date(row[0].value)
        if not dt:
            continue
        conteo_todas[(dt.year, dt.month)] += 1
        if not is_row_yellow(ws_imp, row[0].row):
            conteo[(dt.year, dt.month)] += 1
    if not conteo:
        conteo = conteo_todas
    if not conteo:
        return None, None
    (year, month), _ = conteo.most_common(1)[0]
    return year, month


def detectar_columnas_mes(ws, header_row, year=None, month=None):
    """
    Retorna la columna 'teorico' para el mes/año indicado.
    Si year/month son None o no hay match, retorna la más a la derecha (fallback).
    """
    mes_nombre = _norm(MESES_ES.get(month, '')) if month else None
    yr_str = str(year % 100) if year else None       # "26" para 2026
    yr_full = str(year) if year else None             # "2026"

    teo_col_match = None
    teo_col_fallback = None
    for c in range(1, ws.max_column + 1):
        h = ws.cell(header_row, c).value
        if not h:
            continue
        h_norm = _norm(str(h))
        if 'teorico' not in h_norm:
            continue
        teo_col_fallback = c
        if mes_nombre and yr_str and mes_nombre in h_norm:
            if yr_str in h_norm or (yr_full and yr_full in h_norm):
                teo_col_match = c
    return teo_col_match if teo_col_match is not None else teo_col_fallback


# Tokens (nombre completo + abreviaturas) para reconocer el mes en un header.
_MES_TOKENS = {
    1: ('enero', 'ene'),      2: ('febrero', 'feb'),   3: ('marzo', 'mar'),
    4: ('abril', 'abr'),      5: ('mayo', 'may'),      6: ('junio', 'jun'),
    7: ('julio', 'jul'),      8: ('agosto', 'agost', 'ago'),
    9: ('septiembre', 'setiembre', 'sept', 'set', 'sep'),
    10: ('octubre', 'oct'),   11: ('noviembre', 'nov'), 12: ('diciembre', 'dic'),
}


def mapa_meses_columnas(ws, header_row):
    """Retorna {(year, month): teo_col} parseando los headers 'teorico'.

    Permite imputar cada transferencia en la columna del mes de SU fecha
    (las hojas USD arrastran transferencias de varios meses en la misma hoja).
    """
    mapa = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(header_row, c).value
        if not h:
            continue
        hn = _norm(str(h))
        if 'teorico' not in hn:
            continue
        yr = re.search(r'\b(20\d{2}|\d{2})\b', hn)
        if not yr:
            continue
        y = int(yr.group(1))
        year = y if y > 99 else 2000 + y
        # Matchear por PALABRA (prefijo), no substring: 'ago' no debe matchear
        # dentro de 'pago', ni 'dic' dentro de otra palabra.
        palabras = re.findall(r'[a-z]+', hn)
        mes = None
        for num, toks in _MES_TOKENS.items():
            if any(w.startswith(t) for w in palabras for t in toks):
                mes = num
                break
        if mes is None:
            continue
        mapa[(year, mes)] = c
    return mapa


# Marcador de cuota/lote dentro de col H de imputaciones ("c14", "l13 c29",
# "c8 en ambos lotes mismo monto"). El nombre se corta en el PRIMER marcador:
# lo que sigue son notas sueltas que ensuciaban la búsqueda por nombre.
_RE_MARCA_LOTE_CUOTA = re.compile(r'(?:^|[\s,;.\-])[cl]\s*\d', re.IGNORECASE)


def limpiar_nombre_col_h(col_h_str):
    """Nombre del cliente a partir del texto de col H de imputaciones."""
    m = _RE_MARCA_LOTE_CUOTA.search(col_h_str)
    nombre = col_h_str[:m.start()] if m else col_h_str
    return re.sub(r'\s+', ' ', nombre).strip(' ,;.-')


def cuota_en_col_h(col_h_str):
    """Máxima cuota mencionada en col H ("c8 en ambos lotes" → 8)."""
    nums = []
    for grupo in re.findall(r'(?:^|[\s,;.\-])c\s*(\d+(?:\s*[,yY]\s*\d+)*)', col_h_str):
        nums += [int(n) for n in re.findall(r'\d+', grupo)]
    nums = [n for n in nums if 0 < n <= 200]
    return max(nums) if nums else None


# ── Partes de cuota ("parte de cN" / "completa cN") ─────────────────────────

_RE_PARTE = re.compile(r'parte\s*de\s*(?:la\s*)?(?:cuota\s*)?c?\s*(\d+)', re.IGNORECASE)


def _sin_notas(s):
    """Saca las notas entre paréntesis: 'parte de c3(siguiente transfiere...)'."""
    return re.sub(r'\([^)]*\)?', ' ', s)


def cuota_parcial_de_celda(val):
    """N de 'parte de cN' en una celda de N° de cuota, o None."""
    if not isinstance(val, str):
        return None
    m = _RE_PARTE.search(_sin_notas(val))
    return int(m.group(1)) if m else None


def cuotas_en_celda(val):
    """Cantidad de cuotas (enteras + la parcial) que menciona la celda."""
    if isinstance(val, (int, float)):
        return 1
    if not isinstance(val, str):
        return 0
    return len({int(x) for x in re.findall(r'\d+', _sin_notas(val)) if 0 < int(x) <= 200})


def texto_agregar_cuota(txt, nuevo):
    """Agrega una cuota ('33' o 'parte de c34') a lo que ya dice la celda."""
    if txt is None or (isinstance(txt, str) and not txt.strip()):
        return nuevo
    if isinstance(txt, float) and txt.is_integer():
        txt = int(txt)
    return f'{txt} y {nuevo}'


def texto_completar_cuota(txt, n):
    """'17 y parte de c18' + completar 18 → '17 y 18'; 'parte de c3(nota)' → 3."""
    s = _sin_notas(str(txt))
    s = re.sub(rf'parte\s*de\s*(?:la\s*)?(?:cuota\s*)?c?\s*{n}\b', str(n), s, flags=re.IGNORECASE)
    s = re.sub(r'\s+', ' ', s).strip()
    return int(s) if s.isdigit() else s


def _fmt_num(x):
    x = float(x)
    return str(int(x)) if x.is_integer() else f'{x:.2f}'


def normalizar_telefono(raw):
    """Primer teléfono de la celda en formato wa.me (549 + área + número), o None.

    La col TELEFONO viene en mil formatos: '2995 88-4296', '+54 9 11 6606-2102',
    '2984634923 / 2984967634', '2284 45-9353 (Erica) y 2284 46-5461 (Oscar)'.
    """
    if raw is None:
        return None
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    primero = re.split(r'/|\by\b|\(|,', str(raw))[0]
    d = re.sub(r'\D', '', primero)
    if d.startswith('54'):
        d = d[2:]
    if d.startswith('9') and len(d) == 11:
        d = d[1:]
    if d.startswith('0'):
        d = d[1:]
    return '549' + d if len(d) == 10 else None


def _pesos(x):
    return '$' + f'{round(x):,}'.replace(',', '.')


def mensaje_reclamo(nombre, fecha, monto, n_cuota, saldo, teo, info, otras=None):
    """Texto del reclamo por WhatsApp de una cuota que quedó incompleta."""
    nombre = re.sub(r'\s+', ' ', str(nombre)).strip().title()
    fecha_txt = fecha.strftime('%d/%m') if fecha else ''
    msg = f'Hola {nombre}! Recibimos tu transferencia del {fecha_txt} por {_pesos(monto)}.'
    if info and not info.get('congelado', True):
        msg += (f' Como se hizo después del día 10, la cuota {n_cuota} quedó en {_pesos(teo)}'
                f' (bolsa de {info["kg"]} kg a {_pesos(info["precio_bolsa"])}).')
    elif teo:
        msg += f' La cuota {n_cuota} es de {_pesos(teo)}.'
    msg += f' Para completarla te faltan {_pesos(saldo)}.'
    otras = [(n, sd) for n, sd in (otras or []) if n != n_cuota]
    if otras:
        detalle = ', '.join(f'{_pesos(sd)} de la cuota {n}' for n, sd in otras)
        total = saldo + sum(sd for _, sd in otras)
        msg += f' Además quedaron pendientes {detalle} (total a completar: {_pesos(total)}).'
    msg += ' ¡Muchas gracias!'
    return msg


def link_whatsapp(telefono, mensaje):
    from urllib.parse import quote
    if not telefono:
        return None
    return f'https://wa.me/{telefono}?text={quote(mensaje)}'


# ── Lectura de valores de deudores ───────────────────────────────────────────

_RE_REF_CELDA = re.compile(r'\$?\b([A-Z]{1,3})\$?(\d+)\b')
_RE_ARITMETICA = re.compile(r'^[\d.\s+\-*/()eE]+$')


class LectorDeudores:
    """Valor numérico de una celda de deudores.

    Usa el valor cacheado por Excel; si no está (el archivo lo guardó openpyxl,
    p. ej. la salida de una corrida anterior), evalúa la fórmula si es
    aritmética simple con referencias (=U7*EX$2, =S7/T7, =497000+6000).
    Fórmulas con funciones (SUM, IF...) o referencias rotas devuelven None.
    """

    def __init__(self, wb_valores, wb_formulas):
        self.wb_v = wb_valores
        self.wb_f = wb_formulas
        self._memo = {}

    def formula(self, sname, row, col):
        return self.wb_f[sname].cell(row, col).value

    def valor(self, sname, row, col, _prof=0):
        key = (sname, row, col)
        if key in self._memo:
            return self._memo[key]
        v = self.wb_v[sname].cell(row, col).value
        if not (isinstance(v, (int, float)) and not isinstance(v, bool)):
            f = self.formula(sname, row, col)
            if isinstance(f, (int, float)) and not isinstance(f, bool):
                v = f
            elif isinstance(f, str) and f.startswith('=') and _prof < 10:
                v = self._evaluar(sname, f[1:], _prof)
            else:
                v = None
        self._memo[key] = v
        return v

    def _evaluar(self, sname, cuerpo, prof):
        from openpyxl.utils import column_index_from_string
        faltan = []

        def _ref(m):
            val = self.valor(sname, int(m.group(2)), column_index_from_string(m.group(1)), prof + 1)
            if val is None:
                faltan.append(m.group(0))
                return '0'
            return repr(float(val))

        expr = _RE_REF_CELDA.sub(_ref, cuerpo)
        if faltan or not _RE_ARITMETICA.match(expr):
            return None
        try:
            return float(eval(expr, {'__builtins__': {}}, {}))
        except Exception:
            return None


# ── BOLSA CEMENTO: precio por tramo de fecha ─────────────────────────────────

HOJA_BOLSA = 'BOLSA CEMENTO'
UMBRAL_CUOTA_DEL_MES = 0.9   # pago ≥ 90% del teórico = la cuota del mes pagada de menos (no completa partes viejas)
# El precio se congela hasta el 10, pero una transferencia del 10 que cae en
# finde/feriado aparece acreditada el 12 o 13: hasta el 13 se toma congelado.
DIA_TOPE_CONGELADO = 13
MESES_PARTE = 3   # meses hacia atrás donde se buscan partes de cuota pendientes
FIRMA_BOLSA_50K_HASTA = datetime.datetime(2025, 8, 31)   # contratos con bolsa de 50 kg


def _header_es_del_mes(hn, year, month):
    palabras = re.findall(r'[a-z]+', hn)
    if not any(w.startswith(t) for w in palabras for t in _MES_TOKENS[month]):
        return False
    return str(year) in hn or re.search(rf'\b{year % 100:02d}\b', hn) is not None


def detectar_tramos_bolsa(lector, ws_f, year, month):
    """Precios de la bolsa (25 kg) del mes según el día de pago.

    Lee los encabezados de la fila 1 ('B CEMENTO LOMA NEGRA 25 K 1 AL 10 DE
    SEPT 2026', '... 10 al 15 ...', '... desde el 16 ...') y el precio de la
    fila 2. Ignora las columnas 'los q firmaron hasta agosto 2025' (son el
    mismo precio ×2). Devuelve {'tramos': [(dia_hasta, precio)], 'desc': coef}
    o None si el mes no tiene tramos cargados.
    """
    tramos, desc = [], 1.0
    for c in range(1, ws_f.max_column + 1):
        h = ws_f.cell(1, c).value
        if not isinstance(h, str):
            continue
        hn = _norm(h)
        if not _header_es_del_mes(hn, year, month):
            continue
        if 'desc' in hn:
            v = lector.valor(HOJA_BOLSA, 2, c)
            if v and 0 < v < 1:
                desc = v
            continue
        if 'firmaron' in hn:
            continue
        m = re.search(r'desde\s+el\s+(\d+)', hn)
        if m:
            hasta = 31
        else:
            m = re.search(r'(\d+)\s+al\s+(\d+)', hn)
            if not m:
                continue
            hasta = int(m.group(2))
        precio = lector.valor(HOJA_BOLSA, 2, c)
        if precio:
            tramos.append((hasta, precio))
    if not tramos:
        return None
    return {'tramos': sorted(tramos), 'desc': desc}


def precio_tramo(tramos_info, dia):
    """(precio 25 kg, texto del tramo) que corresponde al día de pago.

    El primer tramo (precio congelado) se estira hasta DIA_TOPE_CONGELADO."""
    primero_hasta, primero_precio = tramos_info['tramos'][0]
    if dia <= max(primero_hasta, DIA_TOPE_CONGELADO):
        return primero_precio, f'congelado (hasta el {primero_hasta})'
    desde = 1
    for hasta, precio in tramos_info['tramos']:
        if dia <= hasta:
            txt = f'del {desde} al {hasta}' if hasta < 31 else f'desde el {desde}'
            return precio, txt
        desde = hasta + 1
    hasta, precio = tramos_info['tramos'][-1]
    return precio, f'desde el {desde}'


def _mismo_cliente(matches):
    """True si todas las filas candidatas son del mismo titular (varios lotes)."""
    return len({_norm(str(n)).strip() for _, _, n in matches}) == 1


def _ultima_col_cuota(ws, row, hist_cols, excluir_col=None):
    """Columna 'NUMERO DE CUOTA' más a la derecha con valor en esa fila.

    Proxy de "hace cuánto se paga este lote": se usa para elegir a qué lotes va
    el reparto cuando un CUIT figura en más lotes de los que cubre el monto
    (p. ej. como cofirmante de un lote de otro titular que no paga hace meses).
    """
    ultima = -1
    for c in hist_cols:
        if c == excluir_col:
            continue
        if max_cuota_celda(ws.cell(row, c).value) is not None:
            ultima = max(ultima, c)
    return ultima


def _col_por_header(ws, header_row, keywords):
    """Primera columna cuyo header (normalizado) contiene alguna de keywords."""
    for c in range(1, ws.max_column + 1):
        h = ws.cell(header_row, c).value
        if not h:
            continue
        hn = _norm(str(h))
        if any(k in hn for k in keywords):
            return c
    return None


def build_sheets_cfg(wb_deu_data, sheets_base, year=None, month=None):
    """Auto-detecta columnas del mes y construye sheets_cfg completo."""
    sheets_cfg = {}
    mes_info = {}
    for sheet_name, base in sheets_base.items():
        try:
            ws = wb_deu_data[sheet_name]
        except KeyError:
            continue
        teo_col = detectar_columnas_mes(ws, base['header_row'], year=year, month=month)
        if teo_col is None:
            continue
        cfg = dict(base)
        # Auto-detectar CUIT/Nombre/LOTE por header (los índices fijos se rompen
        # si el archivo de deudores agrega/quita columnas). Se cae al valor base
        # si no se encuentra el header.
        hr = base['header_row']
        cuit_c   = _col_por_header(ws, hr, ('cuit',))
        nombre_c = _col_por_header(ws, hr, ('nombre',))
        lote_c   = _col_por_header(ws, hr, ('lote',))
        if cuit_c:
            cfg['cuit_col'] = cuit_c
        if nombre_c:
            cfg['nombre_col'] = nombre_c
        if lote_c:
            cfg['lote_col'] = lote_c
        cfg['teo_col'] = teo_col
        cfg['real_col'] = teo_col + 1
        cfg['cuota_col'] = teo_col + 2
        cfg['fecha_col'] = teo_col + 3
        cfg['max_cuota_col'] = 1
        sheets_cfg[sheet_name] = cfg
        header_val = ws.cell(base['header_row'], teo_col).value
        mes_info[sheet_name] = str(header_val).strip() if header_val else f'col {teo_col}'
    return sheets_cfg, mes_info


def build_indices(wb_deu_data, sheets_cfg):
    cuit_index = {}
    nombre_index = {}
    cuota_history_cols = {}

    for sheet_name, cfg in sheets_cfg.items():
        ws_data = wb_deu_data[sheet_name]
        max_row = ultima_fila_con_datos(ws_data, (cfg['nombre_col'], cfg['cuit_col']))

        for r in range(cfg['data_start'], max_row + 1):
            nombre = ws_data.cell(r, cfg['nombre_col']).value
            cuit_raw = ws_data.cell(r, cfg['cuit_col']).value
            if not nombre and not cuit_raw:
                continue
            for c in normalize_cuits(cuit_raw):
                cuit_index.setdefault(c, []).append((sheet_name, r, nombre))

        for r in range(cfg['data_start'], max_row + 1):
            nombre = ws_data.cell(r, cfg['nombre_col']).value
            if not nombre:
                continue
            for palabra in str(nombre).upper().split():
                if len(palabra) >= 4:
                    # normalizar tildes para que "ÁNGEL" y "ANGEL" sean la misma clave
                    nombre_index.setdefault(_norm(palabra).upper(), []).append((sheet_name, r, nombre))

        cols = []
        for c in range(1, ws_data.max_column + 1):
            h = ws_data.cell(cfg['header_row'], c).value
            if h and ('numero' in str(h).lower() or 'n°' in str(h).lower() or 'nro' in str(h).lower()) and 'cuota' in str(h).lower():
                cols.append(c)
        # Guarda TODAS las columnas de N° cuota; el loop excluye la del mes
        # objetivo (que varía por fila cuando cada transferencia va a su propio mes).
        cuota_history_cols[sheet_name] = cols

    return cuit_index, nombre_index, cuota_history_cols


def buscar_en_deudores_por_nombre(nombre_str, nombre_index):
    from collections import Counter

    # Si tiene "/" (ej: "Cerda Gabriel / Iboldi Yanina"), buscar cada parte por separado
    if '/' in nombre_str:
        resultado, seen = [], set()
        for parte in nombre_str.split('/'):
            for item in buscar_en_deudores_por_nombre(parte.strip(), nombre_index):
                k = (item[0], item[1])
                if k not in seen:
                    resultado.append(item)
                    seen.add(k)
        return resultado

    # Normalizar tildes en las palabras buscadas (igual que en el índice)
    palabras = [_norm(p).upper() for p in nombre_str.split() if len(p) >= 4]
    if not palabras:
        return []
    sets = [set((s, r) for s, r, _ in nombre_index.get(p, [])) for p in palabras]
    if not sets:
        return []
    comunes = sets[0]
    for s in sets[1:]:
        comunes &= s
    if not comunes:
        cnt = Counter()
        for p in palabras:
            for s, r, _ in nombre_index.get(p, []):
                cnt[(s, r)] += 1
        max_hits = max(cnt.values()) if cnt else 0
        comunes = {k for k, v in cnt.items() if v == max_hits and v >= max(2, len(palabras) - 1)}
    resultado = []
    seen = set()
    for p in palabras:
        for s, r, n in nombre_index.get(p, []):
            if (s, r) in comunes and (s, r) not in seen:
                resultado.append((s, r, n))
                seen.add((s, r))
    return resultado


# Cuántas hojas anteriores se miran para sacar el nombre de un CUIT/cuenta que
# no está en deudores. Muchos clientes transfieren desde una cuenta (no desde su
# CUIT) y pagan cada 2-3 meses: con una ventana corta el antecedente queda afuera.
HOJAS_PREVIAS_PESOS = 20
HOJAS_PREVIAS_USD = 6


def build_previo(wb_imp, imp_sheet, es_usd=False):
    all_sheets = wb_imp.sheetnames
    try:
        idx_actual = all_sheets.index(imp_sheet)
        previas = list(reversed(all_sheets[:idx_actual]))
        if es_usd:
            previas = [h for h in previas if 'usd' in h.lower()][:HOJAS_PREVIAS_USD]
        else:
            previas = previas[:HOJAS_PREVIAS_PESOS]
    except ValueError:
        previas = []

    cuit_to_nombre_previo = {}
    cuit_to_cuota_previo = {}

    def _indexar(ws_prev, solo_amarillas=False):
        fin = ultima_fila_con_datos(ws_prev, (3, 8))  # C = concepto, H = nombre
        if fin < 4:
            return
        for row in ws_prev.iter_rows(min_row=4, max_row=fin, max_col=9):
            concepto = row[2].value if len(row) > 2 else None
            col_h = row[7].value if len(row) > 7 else None
            if not concepto or not col_h:
                continue
            col_h_str = str(col_h).strip()
            if not col_h_str or 'PAGO MENOS' in col_h_str or 'Saldo' in col_h_str:
                continue
            if solo_amarillas and not is_row_yellow(ws_prev, row[0].row):
                continue
            cuit = extract_cuit_from_concepto(concepto)
            if cuit and cuit not in cuit_to_nombre_previo:
                cuota_prev = cuota_en_col_h(col_h_str)
                if cuota_prev is not None:
                    cuit_to_cuota_previo[cuit] = cuota_prev
                nombre_prev = limpiar_nombre_col_h(col_h_str)
                if nombre_prev:
                    cuit_to_nombre_previo[cuit] = nombre_prev

    # La hoja ACTUAL primero (solo filas ya imputadas/amarillas): las hojas USD
    # son acumulativas y arrastran meses de transferencias del mismo cliente más
    # abajo en la misma hoja. Sin esto, el antecedente "estaba más abajo" nunca
    # se miraba. Va primero porque es el dato más reciente (gana el primero).
    try:
        _indexar(wb_imp[imp_sheet], solo_amarillas=True)
    except Exception:
        pass

    for hoja in previas:
        try:
            ws_prev = wb_imp[hoja]
        except Exception:
            continue
        _indexar(ws_prev)

    return cuit_to_nombre_previo, cuit_to_cuota_previo


def procesar(
    imp_bytes, deu_bytes, imp_sheet,
    es_usd=False,
    tolerance=None,
    max_row=500,
    cuota_override=None,
    comprobantes_cache=None,
    mep_rates=None,
    log_fn=None,
    solo_mes=None,
    forzar_col_mes=None,
    reclamos_out=None,
    reprocesar_pago_menos=False,
    telefonos_extra=None,
    tramos_bolsa_override=None,
):
    """
    Corre la imputación en modo simulación.

    Pesos — pagos incompletos (ver CLAUDE.md, "Partes de cuota"):
      * paga menos que el teórico por más de la tolerancia → "parte de cN"
        (antes: "PAGO MENOS"). Si tiene una parte pendiente, primero la completa.
      * paga la cuota entera → cuota siguiente, sin mirar partes pendientes.
      * paga de más y el sobrante alcanza para completar una parte pendiente →
        "cX y completa cN".
    reclamos_out: lista opcional donde se agregan los reclamos de WhatsApp
    (solo BOLSA CEMENTO) de las cuotas que quedan incompletas.
    reprocesar_pago_menos: vuelve a procesar filas marcadas "PAGO MENOS".
    telefonos_extra: dict CUIT → teléfono (p. ej. de Supabase); gana sobre la
    col TELEFONO de deudores.
    tramos_bolsa_override: {'tramos': [(dia_hasta, precio_25k), ...], 'desc': 1.0}
    para corregir los precios del mes si los encabezados no se pueden leer.

    solo_mes: tupla (year, month) opcional. Si se pasa, se procesan SOLO las
    transferencias fechadas en ese mes y se escribe en la columna de ese mes.
    Sirve para semanas con cambio de mes en el medio (correr una pasada por mes).
    Si es None, autodetecta el mes más frecuente (comportamiento histórico).

    forzar_col_mes: tupla (year, month) opcional. Fuerza la COLUMNA del mes a
    escribir, pero SIN filtrar filas por fecha (procesa todas). Útil para la
    2da pasada de una semana con cambio de mes: "todo lo que no se imputó al mes
    previo va al mes nuevo" (lógica de próxima cuota impaga). Ignorado si se pasa
    solo_mes.
    Retorna (results, pago_menos, pago_mas, ambiguous, sin_fila, usd_en_pesos, mes_info, sheets_cfg).

    usd_en_pesos: clientes de la hoja '$  USD fijo' que pagaron en pesos
    (CUIT no está en las hojas de pesos pero sí en USD fijo). Se convierte el
    monto por el dólar MEP venta del día (mep_rates: dict 'YYYY-MM-DD' -> venta)
    y se reporta para que el usuario los habilite uno por uno.
    """
    if tolerance is None:
        tolerance = 5 if es_usd else 3000
    if cuota_override is None:
        cuota_override = {}
    if comprobantes_cache is None:
        comprobantes_cache = {}
    if log_fn is None:
        log_fn = lambda msg: None

    sheets_base = SHEETS_BASE_USD if es_usd else SHEETS_BASE_PESOS

    # Solo data_only para leer — no cargamos el workbook editable acá
    wb_deu_data = openpyxl.load_workbook(io.BytesIO(deu_bytes), data_only=True)
    wb_deu_f = openpyxl.load_workbook(io.BytesIO(deu_bytes))  # fórmulas crudas
    wb_imp = openpyxl.load_workbook(io.BytesIO(imp_bytes))
    lector = LectorDeudores(wb_deu_data, wb_deu_f)

    if solo_mes:
        tx_year, tx_month = solo_mes
        log_fn(f'Mes forzado (solo_mes): {MESES_ES.get(tx_month, "?")} {tx_year}')
    elif forzar_col_mes:
        tx_year, tx_month = forzar_col_mes
        log_fn(f'Columna forzada (forzar_col_mes): {MESES_ES.get(tx_month, "?")} {tx_year}')
    else:
        tx_year, tx_month = detectar_mes_transferencias(wb_imp[imp_sheet], max_row)
        if tx_year:
            log_fn(f'Mes detectado en transferencias: {MESES_ES.get(tx_month, "?")} {tx_year}')

    sheets_cfg, mes_info = build_sheets_cfg(wb_deu_data, sheets_base, year=tx_year, month=tx_month)
    cuit_index, nombre_index, cuota_history_cols = build_indices(wb_deu_data, sheets_cfg)
    log_fn(f'{len(cuit_index)} CUITs indexados en deudores')

    # En corridas de pesos, indexar también la hoja USD fijo para detectar
    # clientes USD que pagan su cuota en pesos (conversión por dólar MEP)
    usd_cuit_index, usd_hist_cols, usd_cfgs = {}, {}, {}
    if not es_usd:
        usd_cfgs, _ = build_sheets_cfg(wb_deu_data, SHEETS_BASE_USD, year=tx_year, month=tx_month)
        if usd_cfgs:
            usd_cuit_index, _, usd_hist_cols = build_indices(wb_deu_data, usd_cfgs)
            log_fn(f'{len(usd_cuit_index)} CUITs indexados en USD fijo')

    # ── BOLSA CEMENTO: teórico = bolsas/mes × precio del tramo del día de pago
    tramos_bolsa, bolsa_cols = None, {}
    if HOJA_BOLSA in sheets_cfg and tx_year:
        ws_bf = wb_deu_f[HOJA_BOLSA]
        hr_b = sheets_cfg[HOJA_BOLSA]['header_row']
        tramos_bolsa = tramos_bolsa_override or detectar_tramos_bolsa(lector, ws_bf, tx_year, tx_month)
        bolsa_cols = {
            'bolsas_mes': _col_por_header(ws_bf, hr_b, ('bolsas por mes',)),
            'fecha_firma': _col_por_header(ws_bf, hr_b, ('fecha firma',)),
            'firmaron': {c for c in range(1, ws_bf.max_column + 1)
                         if isinstance(ws_bf.cell(1, c).value, str)
                         and 'firmaron' in _norm(ws_bf.cell(1, c).value)},
        }
        if tramos_bolsa:
            txt = ', '.join(f'hasta el {h}: ${p:,.0f}' for h, p in tramos_bolsa['tramos'])
            log_fn(f'Bolsa cemento {MESES_ES[tx_month]}: {txt} (bolsa 25 kg; ×2 si firmó hasta ago-2025)')
        else:
            log_fn('Bolsa cemento: no se encontraron los precios por tramo del mes → se usa el teórico de la planilla')

    def bolsa_doble(srow, teo_col):
        """True si el contrato es de bolsa de 50 kg (firmó hasta agosto 2025).

        Primero mira a qué columna de precio apunta la fórmula del teórico (lo
        que decidió la oficina); si no se puede, la fecha de firma."""
        from openpyxl.utils import column_index_from_string
        f = lector.formula(HOJA_BOLSA, srow, teo_col)
        if isinstance(f, str) and f.startswith('='):
            refs = {column_index_from_string(m.group(1))
                    for m in _RE_REF_CELDA.finditer(f) if m.group(2) == '2'}
            if refs:
                return bool(refs & bolsa_cols['firmaron'])
        fc = bolsa_cols.get('fecha_firma')
        ff = parse_date(wb_deu_data[HOJA_BOLSA].cell(srow, fc).value) if fc else None
        return bool(ff and ff <= FIRMA_BOLSA_50K_HASTA)

    def teo_de(sname, srow, teo_col, fecha):
        """(teórico, info) de una fila. En BOLSA CEMENTO y el mes en curso se
        calcula con el precio del tramo de `fecha`; si no, el de la planilla."""
        if (sname == HOJA_BOLSA and tramos_bolsa and bolsa_cols.get('bolsas_mes')
                and teo_col == sheets_cfg[HOJA_BOLSA]['teo_col']):
            bolsas = lector.valor(sname, srow, bolsa_cols['bolsas_mes'])
            if bolsas:
                if fecha is None:
                    dia = 31
                elif (fecha.year, fecha.month) == (tx_year, tx_month):
                    dia = fecha.day
                else:
                    dia = 1 if (fecha.year, fecha.month) < (tx_year, tx_month) else 31
                precio, tramo = precio_tramo(tramos_bolsa, dia)
                doble = bolsa_doble(srow, teo_col)
                precio_bolsa = precio * (2 if doble else 1)
                teo = round(bolsas * precio_bolsa * tramos_bolsa['desc'])
                return teo, {'dia': dia, 'tramo': tramo,
                             'congelado': precio == tramos_bolsa['tramos'][0][1], 'precio_bolsa': precio_bolsa,
                             'kg': 50 if doble else 25, 'bolsas': bolsas}
        return lector.valor(sname, srow, teo_col), None

    # ── Bloques (teórico, real, N° cuota, fecha) que toca esta corrida
    estado_bloques = {}

    def bloque(sname, srow, cuota_col):
        k = (sname, srow, cuota_col)
        if k not in estado_bloques:
            txt = wb_deu_data[sname].cell(srow, cuota_col).value
            estado_bloques[k] = {
                'base': wb_deu_f[sname].cell(srow, cuota_col - 1).value,
                'sumas': [],
                'txt': txt,
                'fecha': parse_date(wb_deu_data[sname].cell(srow, cuota_col + 1).value),
                'fecha_nueva': None,
            }
        return estado_bloques[k]

    def bloque_vacio(b):
        return b['base'] in (None, '') and not b['sumas'] and b['txt'] in (None, '')

    def pagado(sname, srow, cuota_col):
        b = bloque(sname, srow, cuota_col)
        base = lector.valor(sname, srow, cuota_col - 1) if b['base'] not in (None, '') else 0
        return (base or 0) + sum(b['sumas'])

    def pendientes(sname, srow):
        """Cuotas 'parte de cN' con saldo, de los últimos MESES_PARTE meses
        (más vieja primero). El saldo se calcula con el teórico del mes de la
        parte (en BOLSA, al precio del día del primer pago)."""
        cfg_p = sheets_cfg[sname]
        cols = [c for c in cuota_history_cols.get(sname, []) if c <= cfg_p['cuota_col']][-MESES_PARTE:]
        out = []
        for c in cols:
            b = bloque(sname, srow, c)
            n = cuota_parcial_de_celda(b['txt'])
            if n is None:
                continue
            teo_p, info_p = teo_de(sname, srow, c - 2, b['fecha'])
            if not teo_p:
                continue
            saldo = cuotas_en_celda(b['txt']) * teo_p - pagado(sname, srow, c)
            if saldo > tolerance:
                out.append({'col': c, 'n': n, 'saldo': round(saldo), 'teo': teo_p, 'info': info_p})
        return out

    def max_cuota_fila(sname, srow):
        mx = None
        for hc in cuota_history_cols.get(sname, []):
            k = (sname, srow, hc)
            val = estado_bloques[k]['txt'] if k in estado_bloques else wb_deu_data[sname].cell(srow, hc).value
            n = max_cuota_celda(val)
            if n is not None and (mx is None or n > mx):
                mx = n
        return mx

    def snapshot(sname, srow, cuota_col):
        """Valores finales a escribir en un bloque tocado."""
        b = estado_bloques[(sname, srow, cuota_col)]
        partes = []
        base = b['base']
        if isinstance(base, str) and base.startswith('='):
            partes.append(base[1:])
        elif isinstance(base, (int, float)):
            partes.append(_fmt_num(base))
        partes += [_fmt_num(x) for x in b['sumas']]
        if len(partes) == 1 and not isinstance(base, str):
            real_val = b['sumas'][0] if b['sumas'] else base
        else:
            real_val = '=' + '+'.join(partes)
        txt = b['txt']
        max_ent = None
        if isinstance(txt, str):
            n_parc = cuota_parcial_de_celda(txt)
            enteras = [int(x) for x in re.findall(r'\d+', _sin_notas(txt)) if int(x) != n_parc]
            max_ent = max(enteras) if enteras else None
        return {'hoja': sname, 'fila': srow, 'real_col': cuota_col - 1, 'cuota_col': cuota_col,
                'fecha_col': cuota_col + 1, 'real': real_val, 'cuota': txt,
                'fecha': b['fecha_nueva'], 'max_cuota': max_ent}

    col_tel = {}
    for s_n, s_cfg in sheets_cfg.items():
        col_tel[s_n] = _col_por_header(wb_deu_f[s_n], s_cfg['header_row'], ('telefono',))

    cuit_to_nombre_previo, cuit_to_cuota_previo = build_previo(wb_imp, imp_sheet, es_usd)
    log_fn(f'{len(cuit_to_nombre_previo)} CUITs con nombre desde hojas anteriores')
    log_fn(f'{len(comprobantes_cache)} CUITs en cache de comprobantes')

    def _buscar_nombre(nombre_str):
        return buscar_en_deudores_por_nombre(nombre_str, nombre_index)

    # Mapa mes→columna por hoja: cada transferencia se imputa en la columna del
    # mes de SU fecha. SOLO en USD, porque su hoja acumula transferencias de
    # varios meses. En pesos cada hoja "S NNN" es de una sola semana/mes, así que
    # se usa siempre la columna del mes detectado (comportamiento previo estable;
    # evita que un teórico duplicado/con fórmula del deudores rompa la lectura).
    # solo_mes/forzar_col_mes fuerzan siempre la columna (flujo de mes partido).
    col_maps = {}
    ruteo_por_mes = es_usd and not (solo_mes or forzar_col_mes)
    if ruteo_por_mes:
        for sname in list(sheets_cfg) + list(usd_cfgs):
            try:
                hr = (sheets_cfg.get(sname) or usd_cfgs.get(sname))['header_row']
                col_maps[sname] = mapa_meses_columnas(wb_deu_data[sname], hr)
            except (KeyError, TypeError):
                col_maps[sname] = {}

    def cols_de(sname, fecha_dt, base_cfg):
        """Devuelve la cfg de columnas del mes de la transferencia (o la base)."""
        if not ruteo_por_mes or fecha_dt is None:
            return base_cfg
        teo = col_maps.get(sname, {}).get((fecha_dt.year, fecha_dt.month))
        if teo is None:
            return base_cfg
        cfg = dict(base_cfg)
        cfg['teo_col'] = teo
        cfg['real_col'] = teo + 1
        cfg['cuota_col'] = teo + 2
        cfg['fecha_col'] = teo + 3
        return cfg

    ws_imp = wb_imp[imp_sheet]
    results    = []
    ambiguous  = []
    pago_menos = []
    pago_mas   = []
    sin_fila   = []  # nombre conocido pero sin fila en deudores → prellena col H sin amarillo
    usd_en_pesos = []  # clientes USD fijo que pagaron en pesos (conversión MEP)
    written_deu_rows = {}  # (sname, srow) -> {'ident': str, 'last_cuota': int}

    EXCESO_LIMITE   = 50 if es_usd else 50_000   # exceso máximo para imputar normalmente
    MULTI_TOL_RATIO = 0.05                        # tolerancia proporcional para múltiplos

    for row in ws_imp.iter_rows(min_row=4, max_row=max_row):
        fecha_val = row[0].value
        monto_val = row[5].value
        concepto = row[2].value
        col_h_val = row[7].value

        if fecha_val is None and monto_val is None:
            continue

        row_num = row[0].row
        fecha_dt = parse_date(fecha_val)

        if is_row_yellow(ws_imp, row_num):
            continue

        # Semana con cambio de mes: procesar solo las filas del mes pedido.
        if solo_mes:
            if not (fecha_dt and fecha_dt.year == solo_mes[0] and fecha_dt.month == solo_mes[1]):
                continue

        if col_h_val and 'Saldo Disponible' in str(col_h_val):
            continue
        if col_h_val and 'PAGO MENOS' in str(col_h_val) and not reprocesar_pago_menos:
            continue

        cuit_raw = extract_cuit_from_concepto(concepto)
        if not cuit_raw:
            ambiguous.append({'row': row_num, 'motivo': 'Sin CUIT extraíble', 'concepto': str(concepto)[:60] if concepto else '', 'monto': monto_val, 'fecha': fecha_val})
            continue

        matches = cuit_index.get(cuit_raw, [])

        # Cliente de USD fijo pagando en pesos: el CUIT no está en las hojas de
        # pesos pero sí en la hoja USD. Convertir por MEP y reportar aparte.
        if not matches and cuit_raw in usd_cuit_index:
            umatches = usd_cuit_index[cuit_raw]
            # elegir la primera fila sin imputar este mes (real y cuota vacíos)
            destino = None
            for (u_sname, u_srow, u_snombre) in umatches:
                u_cfg = cols_de(u_sname, fecha_dt, usd_cfgs[u_sname])
                u_real = wb_deu_data[u_sname].cell(u_srow, u_cfg['real_col']).value
                u_cuota_act = wb_deu_data[u_sname].cell(u_srow, u_cfg['cuota_col']).value
                ya_usada = (u_sname, u_srow) in written_deu_rows
                if u_real is None and not isinstance(u_cuota_act, (int, float)) and not ya_usada:
                    destino = (u_sname, u_srow, u_snombre, u_cuota_act, u_cfg)
                    break
            if destino is None:
                ambiguous.append({'row': row_num, 'motivo': f'CUIT {cuit_raw} en USD fijo pero el mes ya está imputado', 'cliente': umatches[0][2], 'cuit': cuit_raw, 'monto': monto_val})
                continue

            u_sname, u_srow, u_snombre, u_cuota_act, u_cfg = destino
            tasa, tasa_fecha = mep_para_fecha(mep_rates or {}, fecha_dt) if fecha_dt else (None, None)

            monto_num = monto_val if isinstance(monto_val, (int, float)) else 0
            teo_usd = wb_deu_data[u_sname].cell(u_srow, u_cfg['teo_col']).value
            teo_usd = teo_usd if isinstance(teo_usd, (int, float)) else None
            equiv = round(monto_num / tasa, 2) if (tasa and monto_num) else None
            dif = round(equiv - teo_usd, 2) if (equiv is not None and teo_usd is not None) else None

            # número de cuota: misma lógica que el flujo normal, sobre la fila USD
            if isinstance(u_cuota_act, str) and 'parte' in u_cuota_act.lower():
                m = re.search(r'\d+', u_cuota_act)
                next_cuota = int(m.group()) + 1 if m else None
            else:
                max_hist = None
                for hc in usd_hist_cols.get(u_sname, []):
                    if hc == u_cfg['cuota_col']:
                        continue
                    n_celda = max_cuota_celda(wb_deu_data[u_sname].cell(u_srow, hc).value)
                    if n_celda is not None and (max_hist is None or n_celda > max_hist):
                        max_hist = n_celda
                next_cuota = max_hist + 1 if max_hist is not None else None

            if next_cuota is None:
                ambiguous.append({'row': row_num, 'motivo': f'CUIT {cuit_raw} en USD fijo (pagó en pesos) pero no se pudo determinar cuota', 'cliente': u_snombre, 'cuit': cuit_raw, 'monto': monto_val, 'hoja': u_sname, 'hoja_fila': u_srow})
                continue

            usd_en_pesos.append({
                'imp_row': row_num,
                'cuit': cuit_raw,
                'cliente': u_snombre,
                'hoja': u_sname,
                'hoja_fila': u_srow,
                'fecha': fecha_dt,
                'monto_pesos': monto_num,
                'mep': tasa,
                'mep_fecha': tasa_fecha,
                'equiv_usd': equiv,
                'teo_usd': teo_usd,
                'dif_usd': dif,
                'cuota': next_cuota,
                'real_col': u_cfg['real_col'],
                'cuota_col': u_cfg['cuota_col'],
                'fecha_col': u_cfg['fecha_col'],
            })
            written_deu_rows[(u_sname, u_srow)] = {'ident': cuit_raw, 'last_cuota': next_cuota}
            continue

        # Identidad del cliente para no apilar cuotas de dos personas distintas
        # en la misma fila de deudores. Es el CUIT cuando el match vino del
        # índice; si vino por nombre, es el nombre (un mismo cliente puede
        # transferir desde varias cuentas distintas y ninguna ser su CUIT).
        ident = cuit_raw

        if not matches:
            nombre_previo = cuit_to_nombre_previo.get(cuit_raw)
            if nombre_previo:
                m2 = _buscar_nombre(nombre_previo)
                if len(m2) == 1:
                    matches = m2
                    ident = _norm(nombre_previo)
                    log_fn(f'Fila {row_num}: fallback nombre "{nombre_previo}" → {m2[0][2]}')
                elif len(m2) > 1 and _mismo_cliente(m2):
                    # Varios lotes del MISMO titular: no es ambiguo, es un
                    # cliente con N lotes → sigue por la lógica de reparto.
                    matches = m2
                    ident = _norm(nombre_previo)
                    log_fn(f'Fila {row_num}: fallback nombre "{nombre_previo}" → {m2[0][2]} ({len(m2)} lotes)')
                elif len(m2) > 1:
                    ambiguous.append({'row': row_num, 'motivo': f'CUIT {cuit_raw} no en deudores; nombre "{nombre_previo}" da {len(m2)} candidatos', 'concepto': str(concepto)[:60], 'monto': monto_val, 'fecha': fecha_val, 'matches': [(n, None) for _, _, n in m2]})
                    continue
                else:
                    sin_fila.append({'imp_row': row_num, 'nombre': nombre_previo})
                    ambiguous.append({'row': row_num, 'motivo': f'CUIT {cuit_raw} no en deudores; "{nombre_previo}" no encontrado en deudores', 'concepto': str(concepto)[:60], 'monto': monto_val, 'fecha': fecha_val})
                    continue
            else:
                nombre_comp = comprobantes_cache.get(cuit_raw)
                if nombre_comp:
                    m2 = _buscar_nombre(nombre_comp)
                    if len(m2) == 1:
                        matches = m2
                        ident = _norm(nombre_comp)
                        log_fn(f'Fila {row_num}: fallback comprobantes "{nombre_comp}" → {m2[0][2]}')
                    elif len(m2) > 1 and _mismo_cliente(m2):
                        matches = m2
                        ident = _norm(nombre_comp)
                        log_fn(f'Fila {row_num}: fallback comprobantes "{nombre_comp}" → {m2[0][2]} ({len(m2)} lotes)')
                    elif len(m2) > 1:
                        ambiguous.append({'row': row_num, 'motivo': f'CUIT {cuit_raw} en comprobantes como "{nombre_comp}"; da {len(m2)} candidatos', 'concepto': str(concepto)[:60], 'monto': monto_val, 'fecha': fecha_val, 'matches': [(n, None) for _, _, n in m2]})
                        continue
                    else:
                        sin_fila.append({'imp_row': row_num, 'nombre': nombre_comp})
                        ambiguous.append({'row': row_num, 'motivo': f'CUIT {cuit_raw} en comprobantes como "{nombre_comp}"; no encontrado en deudores', 'concepto': str(concepto)[:60], 'monto': monto_val, 'fecha': fecha_val})
                        continue
                else:
                    ambiguous.append({'row': row_num, 'motivo': f'CUIT {cuit_raw} no encontrado en deudores, semanas anteriores ni comprobantes', 'concepto': str(concepto)[:60], 'monto': monto_val, 'fecha': fecha_val})
                    continue

        # monto numérico para todas las decisiones de abajo
        monto_num = monto_val if isinstance(monto_val, (int, float)) else 0

        # reparto_lotes=True → una cuota a cada lote (cada uno con su teórico)
        reparto_lotes = False
        if len(matches) > 1:
            candidatos = []
            for (sname, srow, snombre) in matches:
                cfg = cols_de(sname, fecha_dt, sheets_cfg[sname])
                ws_d = wb_deu_data[sname]
                teo, _ = teo_de(sname, srow, cfg['teo_col'], fecha_dt)
                real = wb_deu_f[sname].cell(srow, cfg['real_col']).value
                # Un lote con "parte de cN" en el mes sigue disponible (en pesos)
                if not es_usd and cuota_parcial_de_celda(ws_d.cell(srow, cfg['cuota_col']).value) is not None:
                    real = None
                candidatos.append((sname, srow, snombre, teo, real))

            sin_imputar = [(s, r, n, t, rv) for s, r, n, t, rv in candidatos if rv is None]
            disponibles = [x for x in sin_imputar if (x[0], x[1]) not in written_deu_rows]
            # Ningún lote libre, pero alguno tiene una parte pendiente (de la
            # planilla o de esta misma corrida): el pago puede ser para
            # completarla (se decide abajo).
            con_pend = [] if (disponibles or es_usd) else [
                (x, pendientes(x[0], x[1])) for x in candidatos]
            con_pend = [(x, pe) for x, pe in con_pend if pe]
            if con_pend:
                x, _pe = min(con_pend, key=lambda xp: abs(xp[1][0]['saldo'] - monto_num))
                targets = [x]
            elif not sin_imputar:
                ambiguous.append({'row': row_num, 'motivo': 'Todos los matches ya tienen el mes imputado', 'cuit': cuit_raw, 'monto': monto_val, 'matches': [(n, t) for _, _, n, t, _ in candidatos]})
                continue

            if con_pend:
                pass
            elif not disponibles:
                ambiguous.append({'row': row_num, 'motivo': 'Todos los lotes ya asignados en este run', 'cuit': cuit_raw, 'monto': monto_val, 'matches': [(n, t) for _, _, n, t, _ in sin_imputar]})
                continue
            elif len(disponibles) > 1:
                # ¿El monto cubre la SUMA de los teóricos de N lotes sin imputar?
                # → repartir una cuota a cada uno (cada lote con su propio teórico).
                # Se prueba N de mayor a menor: antes solo se repartía si cubría
                # TODOS los lotes, y bastaba un lote de más (p. ej. uno donde el
                # CUIT figura como cofirmante y hace meses no se paga) para que
                # el monto nunca cerrara y las N cuotas terminaran apiladas en un
                # solo lote. Orden de preferencia: lote con pago más reciente.
                def _recencia(x):
                    x_cfg = cols_de(x[0], fecha_dt, sheets_cfg[x[0]])
                    return _ultima_col_cuota(wb_deu_data[x[0]], x[1],
                                             cuota_history_cols.get(x[0], []),
                                             x_cfg['cuota_col'])

                orden = sorted(disponibles, key=lambda x: (-_recencia(x), x[0], x[1]))
                subset = None
                for n in range(len(orden), 1, -1):
                    cand = orden[:n]
                    suma_teo = sum((t or 0) for _, _, _, t, _ in cand)
                    multi_tol = max(tolerance, suma_teo * MULTI_TOL_RATIO)
                    if suma_teo > 0 and abs(monto_num - suma_teo) <= multi_tol:
                        subset = cand
                        break
                if subset:
                    targets = sorted(subset, key=lambda x: (x[0], x[1]))
                    reparto_lotes = True
                else:
                    # No alcanza para ningún reparto (pagó una sola cuota o un
                    # monto raro): imputar al lote cuyo teórico mejor coincide.
                    targets = [min(disponibles, key=lambda x: abs((x[3] or 0) - monto_num))]
            else:
                targets = [disponibles[0]]
        else:
            sname, srow, snombre = matches[0]
            cfg = cols_de(sname, fecha_dt, sheets_cfg[sname])
            ws_d = wb_deu_data[sname]
            teo, _ = teo_de(sname, srow, cfg['teo_col'], fecha_dt)
            real = wb_deu_f[sname].cell(srow, cfg['real_col']).value
            if not es_usd and cuota_parcial_de_celda(ws_d.cell(srow, cfg['cuota_col']).value) is not None:
                real = None
            targets = [(sname, srow, snombre, teo, real)]

        sname, srow, snombre, teo_val, real_existente = targets[0]
        cfg = sheets_cfg[sname]

        # Pesos (un solo lote destino): lógica de partes de cuota
        modo_partes = not es_usd and not reparto_lotes
        if modo_partes:
            b_act = bloque(sname, srow, cfg['cuota_col'])
            ya_imputado = real_existente is not None
            pend = pendientes(sname, srow)
            if ya_imputado and not pend:
                ambiguous.append({'row': row_num, 'motivo': f'Mes ya imputado ({real_existente}) en {sname} fila {srow}', 'cliente': snombre, 'cuit': cuit_raw, 'monto': monto_val})
                continue
        elif real_existente is not None:
            ambiguous.append({'row': row_num, 'motivo': f'Mes ya imputado ({real_existente}) en {sname} fila {srow}', 'cliente': snombre, 'cuit': cuit_raw, 'monto': monto_val})
            continue

        deu_key = (sname, srow)
        prev_assignment = written_deu_rows.get(deu_key)
        if prev_assignment is not None and prev_assignment['ident'] != ident:
            ambiguous.append({'row': row_num, 'motivo': f'Destino duplicado (otro cliente) en {sname} fila {srow}', 'cliente': snombre, 'cuit': cuit_raw, 'monto': monto_val})
            continue

        teo_num = teo_val if isinstance(teo_val, (int, float)) else 0

        # GUARDA CRÍTICA: si el teórico del mes viene vacío o 0 no se puede
        # decidir si pagó de más / de menos, y el monto terminaría imputado como
        # cuota entera sin control. Pasa cuando la columna del mes es una fórmula
        # (=U..*coef) cuyo cache openpyxl lee como None, o el coeficiente del mes
        # todavía no está cargado. En ese caso NO imputar: reportar para revisar.
        if reparto_lotes:
            teos_lote = [t if isinstance(t, (int, float)) else 0
                         for _, _, _, t, _ in targets]
            if any(t <= 0 for t in teos_lote):
                ambiguous.append({'row': row_num, 'motivo': 'Teórico del mes vacío o 0 en algún lote (¿fórmula sin calcular / coeficiente del mes sin cargar?). NO imputado — revisar teórico y reimputar', 'cliente': snombre, 'cuit': cuit_raw, 'monto': monto_val})
                continue
        else:
            if teo_num <= 0:
                ambiguous.append({'row': row_num, 'motivo': f'Teórico del mes vacío o 0 en {sname} fila {srow} (¿fórmula sin calcular / coeficiente del mes sin cargar?). NO imputado — revisar teórico y reimputar', 'cliente': snombre, 'cuit': cuit_raw, 'monto': monto_val})
                continue

        n_cuotas_partes = None
        if modo_partes:
            # ¿Cuota entera, parte de cuota, o completa una parte pendiente?
            #  - entera (±tolerancia) o múltiplo → cuota(s) siguiente(s), sin
            #    mirar partes pendientes (no "roba" plata de la cuota nueva).
            #  - menos → completa partes pendientes (más vieja primero) y lo que
            #    sobra queda como "parte de" la cuota siguiente.
            #  - más → si el sobrante alcanza para completar una parte pendiente,
            #    la completa; si no, imputa normal (o PAGO MAS si es mucho).
            m = monto_num
            dif = m - teo_num
            n_mult = 1
            if dif > tolerance:
                n = round(m / teo_num)
                if n >= 2 and abs(m - n * teo_num) <= max(tolerance, teo_num * MULTI_TOL_RATIO):
                    n_mult = n
            completar, acumular, parte_nueva, enteras = [], None, None, 0
            resto = m
            if abs(dif) <= tolerance or n_mult >= 2:
                enteras = n_mult
            elif dif < 0 and (not pend or (
                    m > sum(p['saldo'] for p in pend) + tolerance
                    and m >= teo_num * UMBRAL_CUOTA_DEL_MES)):
                # Es la cuota del mes pagada de menos (típico: pagó con el
                # precio congelado después del 10): no completa partes viejas.
                # Si el pago no pasa de lo que debe de partes anteriores, sí
                # va a completarlas (rama de abajo).
                parte_nueva = m
            elif dif < 0:
                for p in pend:
                    if resto >= p['saldo'] - tolerance:
                        completar.append(p)
                        resto -= p['saldo']
                    else:
                        acumular = p
                        break
                if acumular is None and resto > tolerance:
                    parte_nueva = resto
            else:
                sobrante = dif
                for p in pend:
                    if sobrante >= p['saldo'] - tolerance:
                        completar.append(p)
                        sobrante -= p['saldo']
                    else:
                        break
                if not completar and dif > EXCESO_LIMITE:
                    pago_mas.append({'row': row_num, 'cliente': snombre, 'cuit': cuit_raw, 'transferido': monto_num, 'teorico': round(teo_num), 'diferencia': round(dif)})
                    continue
                enteras = 1

            if enteras and ya_imputado:
                ambiguous.append({'row': row_num, 'motivo': f'Mes ya imputado ({real_existente}) en {sname} fila {srow}', 'cliente': snombre, 'cuit': cuit_raw, 'monto': monto_val})
                continue

            if not (completar or acumular or parte_nueva is not None) and bloque_vacio(b_act):
                # Cuota(s) entera(s) en un mes vacío: camino de siempre
                n_cuotas_partes = enteras
            else:
                x_max = max_cuota_fila(sname, srow)
                if x_max is None:
                    x_max = cuit_to_cuota_previo.get(cuit_raw)
                if x_max is None and cuit_raw in cuota_override:
                    x_max = cuota_override[cuit_raw] - 1
                if x_max is None and (enteras or parte_nueva is not None):
                    ambiguous.append({'row': row_num, 'motivo': 'No se pudo determinar número de cuota (sin historial)', 'cliente': snombre, 'cuit': cuit_raw, 'hoja': sname, 'hoja_fila': srow})
                    continue
                x_sig = (x_max or 0) + 1

                piezas, tocados, apuntes = [], [], []

                def _anotar(col, monto):
                    b = bloque(sname, srow, col)
                    if bloque_vacio(b):
                        b['fecha_nueva'] = fecha_dt
                        b['fecha'] = fecha_dt
                    b['sumas'].append(monto)
                    apuntes.append(b)
                    if col not in tocados:
                        tocados.append(col)
                    return b

                if enteras:
                    cuotas_ent = list(range(x_sig, x_sig + enteras))
                    b = _anotar(cfg['cuota_col'], m - sum(p['saldo'] for p in completar))
                    for c_n in cuotas_ent:
                        b['txt'] = texto_agregar_cuota(b['txt'], c_n)
                    piezas.append(f'c{_format_cuotas(cuotas_ent)}')
                for p in completar:
                    b = _anotar(p['col'], p['saldo'])
                    b['txt'] = texto_completar_cuota(b['txt'], p['n'])
                    if piezas and piezas[-1].startswith('completa '):
                        piezas[-1] += f" y c{p['n']}"
                    else:
                        piezas.append(f"completa c{p['n']}")
                reclamo = None
                completadas = {p['col'] for p in completar}
                if acumular is not None:
                    _anotar(acumular['col'], resto)
                    piezas.append(f"parte de c{acumular['n']}")
                    reclamo = (acumular['n'], acumular['saldo'] - resto, acumular['teo'], acumular['info'])
                    completadas.add(acumular['col'])
                if parte_nueva is not None:
                    b = _anotar(cfg['cuota_col'], parte_nueva)
                    b['txt'] = texto_agregar_cuota(b['txt'], f'parte de c{x_sig}')
                    piezas.append(f'parte de c{x_sig}')
                    _, info_act = teo_de(sname, srow, cfg['teo_col'], fecha_dt)
                    reclamo = (x_sig, teo_num - parte_nueva, teo_num, info_act)
                # otras partes que siguen abiertas (se mencionan en el reclamo)
                otras_pend = [(p['n'], p['saldo']) for p in pend if p['col'] not in completadas]
                # Redondeo: lo que no se asignó (dentro de la tolerancia) va al
                # último apunte, así la suma de "pago real" da lo transferido.
                if not enteras:
                    resto_sin_asignar = m - sum(p['saldo'] for p in completar) - (resto if acumular is not None else 0) - (parte_nueva or 0)
                    if resto_sin_asignar:
                        apuntes[-1]['sumas'][-1] += resto_sin_asignar

                usar_lote_p = len(matches) > 1
                lote_val = wb_deu_data[sname].cell(srow, cfg['lote_col']).value if usar_lote_p else None
                lote_str_p = f' l{lote_val}' if lote_val is not None else ''
                written_deu_rows[(sname, srow)] = {'ident': ident, 'last_cuota': max_cuota_fila(sname, srow) or x_sig}
                results.append({
                    'imp_row': row_num,
                    'cuit': cuit_raw,
                    'cliente': snombre,
                    'lote_str': lote_str_p,
                    'hoja': sname,
                    'hoja_fila': srow,
                    'monto_real': m,
                    'monto_teo': round(teo_num),
                    'diferencia': round(m - teo_num),
                    'cuota': ' y '.join(piezas),
                    'fecha': fecha_dt,
                    'etiqueta': piezas,
                    'bloques': [snapshot(sname, srow, c) for c in tocados],
                })
                if reclamo and reclamo[1] > tolerance and sname == HOJA_BOLSA and reclamos_out is not None:
                    n_r, saldo_r, teo_r, info_r = reclamo
                    tel_raw = wb_deu_data[sname].cell(srow, col_tel[sname]).value if col_tel.get(sname) else None
                    tel = (telefonos_extra or {}).get(cuit_raw) or normalizar_telefono(tel_raw)
                    msg = mensaje_reclamo(snombre, fecha_dt, m, n_r, saldo_r, teo_r, info_r, otras_pend)
                    reclamos_out.append({
                        'imp_row': row_num, 'cliente': snombre, 'cuit': cuit_raw,
                        'hoja_fila': srow, 'cuota': n_r, 'transferido': m,
                        'teorico': round(teo_r), 'saldo': round(saldo_r),
                        'saldo_total': round(saldo_r + sum(sd for _, sd in otras_pend)),
                        'fecha': fecha_dt, 'tramo': (info_r or {}).get('tramo'),
                        'telefono_planilla': tel_raw, 'telefono': tel,
                        'mensaje': msg,
                        'whatsapp': link_whatsapp(tel, msg),
                    })
                continue

        if reparto_lotes:
            # Una cuota a cada lote; cada lote lleva su propio teórico como pago real.
            usar = targets
            counts = [1] * len(usar)
            montos_por_lote = [t if isinstance(t, (int, float)) else 0
                               for _, _, _, t, _ in usar]
        else:
            if monto_num < teo_num and (teo_num - monto_num) > tolerance:
                pago_menos.append({'row': row_num, 'cliente': snombre, 'cuit': cuit_raw, 'transferido': monto_num, 'teorico': round(teo_num, 2 if es_usd else 0), 'diferencia': round(teo_num - monto_num, 2 if es_usd else 0)})
                continue

            # Exceso positivo: verificar si es múltiplo del teórico o excede el límite.
            # (Solo aplica a UN lote: el reparto entre lotes ya se decidió arriba.)
            n_cuotas = 1
            if n_cuotas_partes is not None:
                n_cuotas = n_cuotas_partes   # ya decidido arriba (pesos)
            elif teo_num > 0 and monto_num > teo_num + tolerance:
                exceso = monto_num - teo_num
                n = round(monto_num / teo_num)
                multi_tol = max(tolerance, teo_num * MULTI_TOL_RATIO)
                if n >= 2 and abs(monto_num - n * teo_num) <= multi_tol:
                    n_cuotas = n
                elif exceso > EXCESO_LIMITE:
                    pago_mas.append({'row': row_num, 'cliente': snombre, 'cuit': cuit_raw, 'transferido': monto_num, 'teorico': round(teo_num, 2 if es_usd else 0), 'diferencia': round(exceso, 2 if es_usd else 0)})
                    continue

            usar = [targets[0]]
            counts = [n_cuotas]
            monto_por_cuota = round(monto_num / n_cuotas, 2 if es_usd else 0)
            montos_por_lote = [monto_por_cuota]

        # Determinar la cuota de cada fila destino sin tocar estado, para poder
        # abortar limpio si alguna falla
        planes = []
        error_motivo = None
        for (p_sname, p_srow, p_snombre, p_teo, _p_real), cnt, monto_lote in zip(usar, counts, montos_por_lote):
            p_cfg = cols_de(p_sname, fecha_dt, sheets_cfg[p_sname])
            prev = written_deu_rows.get((p_sname, p_srow))
            cuota_col_val = wb_deu_data[p_sname].cell(p_srow, p_cfg['cuota_col']).value

            if prev is not None:
                # Mismo cliente, misma fila deudores → cuota siguiente
                next_cuota = prev['last_cuota'] + 1
            elif isinstance(cuota_col_val, str) and 'parte' in cuota_col_val.lower():
                n_parte = max_cuota_celda(cuota_col_val)
                if n_parte is not None:
                    next_cuota = n_parte + 1
                else:
                    error_motivo = f'Cuota dice "parte de..." pero no se pudo extraer número ({cuota_col_val!r})'
                    break
            else:
                max_hist = None
                for hc in cuota_history_cols.get(p_sname, []):
                    if hc == p_cfg['cuota_col']:
                        continue
                    n_celda = max_cuota_celda(wb_deu_data[p_sname].cell(p_srow, hc).value)
                    if n_celda is not None and (max_hist is None or n_celda > max_hist):
                        max_hist = n_celda

                if max_hist is not None:
                    next_cuota = max_hist + 1
                elif cuit_raw in cuit_to_cuota_previo:
                    next_cuota = cuit_to_cuota_previo[cuit_raw] + 1
                    log_fn(f'Fila {row_num}: cuota previa {cuit_to_cuota_previo[cuit_raw]}+1={next_cuota}')
                elif cuit_raw in cuota_override:
                    next_cuota = cuota_override[cuit_raw]
                    log_fn(f'Fila {row_num}: cuota override {next_cuota} para {p_snombre}')
                else:
                    error_motivo = 'No se pudo determinar número de cuota (sin historial)'
                    break
            planes.append((p_sname, p_srow, p_snombre, p_teo, next_cuota, cnt, monto_lote))

        if error_motivo:
            ambiguous.append({'row': row_num, 'motivo': error_motivo, 'cliente': snombre, 'cuit': cuit_raw, 'hoja': sname, 'hoja_fila': srow})
            continue

        # Si el cliente tiene más de un lote en deudores (por CUIT o por nombre),
        # el label lleva siempre el lote ("Nombre l6 c13"), aunque los lotes
        # estén a nombres distintos: sin eso no se sabe a cuál fue la cuota.
        # Y si la transferencia se reparte en varias filas, sí o sí
        # ("Nombre l13 c33 y l14 c23").
        usar_lote = len(matches) > 1 or len(planes) > 1

        for p_sname, p_srow, p_snombre, p_teo, next_cuota, cnt, monto_lote in planes:
            p_cfg = cols_de(p_sname, fecha_dt, sheets_cfg[p_sname])
            p_teo_num = p_teo if isinstance(p_teo, (int, float)) else 0
            if usar_lote:
                lote_val = wb_deu_data[p_sname].cell(p_srow, p_cfg['lote_col']).value
                lote_str = f' l{lote_val}' if lote_val is not None else ''
            else:
                lote_str = ''
            written_deu_rows[(p_sname, p_srow)] = {'ident': ident, 'last_cuota': next_cuota + cnt - 1}
            if not es_usd:
                b_n = bloque(p_sname, p_srow, p_cfg['cuota_col'])
                if bloque_vacio(b_n):
                    b_n['fecha_nueva'] = fecha_dt
                    b_n['fecha'] = fecha_dt
                b_n['sumas'].append(monto_lote * cnt)
                for i in range(cnt):
                    b_n['txt'] = texto_agregar_cuota(b_n['txt'], next_cuota + i)
            for i in range(cnt):
                results.append({
                    'imp_row': row_num,
                    'cuit': cuit_raw,
                    'cliente': p_snombre,
                    'lote_str': lote_str,
                    'hoja': p_sname,
                    'hoja_fila': p_srow,
                    'monto_real': monto_lote,
                    'monto_teo': round(p_teo_num, 2 if es_usd else 0),
                    'diferencia': round(monto_lote - p_teo_num, 2 if es_usd else 0),
                    'cuota': next_cuota + i,
                    'fecha': fecha_dt,
                    'real_col': p_cfg['real_col'],
                    'cuota_col': p_cfg['cuota_col'],
                    'fecha_col': p_cfg['fecha_col'],
                })

    wb_deu_data.close()
    wb_deu_f.close()
    wb_imp.close()

    # incluir la cfg de USD fijo para que aplicar() pueda escribir ahí
    sheets_cfg_out = dict(sheets_cfg)
    sheets_cfg_out.update(usd_cfgs)

    return results, pago_menos, pago_mas, ambiguous, sin_fila, usd_en_pesos, mes_info, sheets_cfg_out


def _format_cuotas(cuotas):
    if len(cuotas) == 1:
        return cuotas[0]
    if len(cuotas) == 2:
        return f'{cuotas[0]} y {cuotas[1]}'
    return ', '.join(str(c) for c in cuotas[:-1]) + f' y {cuotas[-1]}'


def _set_cell(cell, val):
    """Escribe val preservando el number_format de la celda. Excepción: si se
    escribe una fecha en una celda 'General' (típico de una columna de un mes
    nuevo aún sin formatear), se aplica formato de fecha para que no se vea como
    número de serie."""
    fmt = cell.number_format
    cell.value = val
    if isinstance(val, (datetime.datetime, datetime.date)) and fmt in (None, 'General'):
        fmt = 'dd/mm/yyyy'
    cell.number_format = fmt


def aplicar(results, pago_menos, imp_bytes, deu_bytes, imp_sheet, sheets_cfg, pago_mas=None, sin_fila=None, usd_en_pesos=None):
    """Carga los workbooks desde bytes, escribe y retorna (imp_bytes, deu_bytes).

    usd_en_pesos: solo las entradas que el usuario habilitó. Escribe en la hoja
    USD fijo el monto EN PESOS en "pago real", más cuota y fecha.
    """
    from collections import defaultdict
    wb_imp = openpyxl.load_workbook(io.BytesIO(imp_bytes))
    wb_deu_edit = openpyxl.load_workbook(io.BytesIO(deu_bytes))

    ws_edit = wb_imp[imp_sheet]

    for p in pago_menos:
        ws_edit.cell(p['row'], 8).value = 'PAGO MENOS'

    for p in (pago_mas or []):
        ws_edit.cell(p['row'], 8).value = 'PAGO MAS'

    # Prellenar nombre en col H sin amarillo (nombre conocido, fila en deudores no encontrada)
    for s in (sin_fila or []):
        cell = ws_edit.cell(s['imp_row'], 8)
        if not cell.value:  # no pisar si ya tiene algo
            cell.value = s['nombre']

    # Imputaciones: agrupar por imp_row (una transferencia puede cubrir N cuotas)
    # Las de partes de cuota ("parte de c33", "completa c32") traen su label y
    # los valores finales de cada bloque ya calculados: se escriben al final.
    con_partes = [r for r in results if r.get('etiqueta')]
    results = [r for r in results if not r.get('etiqueta')]

    imp_groups = defaultdict(list)
    for r in results:
        imp_groups[r['imp_row']].append(r)

    for row_num, grupo in imp_groups.items():
        filas = defaultdict(list)
        for r in grupo:
            filas[(r['hoja'], r['hoja_fila'])].append(r)
        if len(filas) == 1:
            cuotas = [r['cuota'] for r in grupo]
            nombre_corto = str(grupo[0]['cliente'])[:38]
            label = f"{nombre_corto}{grupo[0]['lote_str']} c{_format_cuotas(cuotas)}"
        else:
            # Transferencia repartida en varios lotes: "Nombre l24 c16 y l25 c16"
            segs = []
            for key in sorted(filas):
                rs = filas[key]
                cuotas_fila = _format_cuotas([r['cuota'] for r in rs])
                if rs[0]['lote_str']:
                    segs.append(f"{rs[0]['lote_str'].strip()} c{cuotas_fila}")
                else:
                    segs.append(f"{str(rs[0]['cliente'])[:25]} c{cuotas_fila}")
            if grupo[0]['lote_str']:
                label = f"{str(grupo[0]['cliente'])[:38]} " + ' y '.join(segs)
            else:
                label = ' y '.join(segs)
        ws_edit.cell(row_num, 8).value = label
        for cell in ws_edit[row_num]:
            cell.fill = YELLOW_FILL

    # Clientes USD fijo que pagaron en pesos (solo los habilitados)
    for e in (usd_en_pesos or []):
        ws_edit.cell(e['imp_row'], 8).value = f"{str(e['cliente'])[:38]} c{e['cuota']}"
        for cell in ws_edit[e['imp_row']]:
            cell.fill = YELLOW_FILL
        cfg = sheets_cfg[e['hoja']]
        ws_deu = wb_deu_edit[e['hoja']]
        for col, val in [
            (e.get('real_col',  cfg['real_col']),  e['monto_pesos']),
            (e.get('cuota_col', cfg['cuota_col']), e['cuota']),
            (e.get('fecha_col', cfg['fecha_col']), e['fecha']),
        ]:
            _set_cell(ws_deu.cell(e['hoja_fila'], col), val)

    # Deudores: agrupar por fila + columna de mes. Una misma fila (lote) puede
    # recibir cuotas de meses distintos en este run (transferencias de junio y
    # julio del mismo cliente) → cada mes va a su propio bloque de columnas.
    deu_groups = defaultdict(list)
    for r in results:
        real_col = r.get('real_col', sheets_cfg[r['hoja']]['real_col'])
        deu_groups[(r['hoja'], r['hoja_fila'], real_col)].append(r)

    for (sname, srow, _rc), grupo in deu_groups.items():
        cfg = sheets_cfg[sname]
        ws_deu = wb_deu_edit[sname]
        cuotas = [r['cuota'] for r in grupo]
        for col, val in [
            (grupo[0].get('real_col',  cfg['real_col']),  sum(r['monto_real'] for r in grupo)),
            (grupo[0].get('cuota_col', cfg['cuota_col']), _format_cuotas(cuotas)),
            (grupo[0].get('fecha_col', cfg['fecha_col']), grupo[0]['fecha']),
        ]:
            _set_cell(ws_deu.cell(srow, col), val)
        # Si se imputan 2+ cuotas en la misma celda (ej "26 y 27"), la fórmula
        # de la columna MAYOR CUOTA no sabe leer ese texto → escribir la cuota
        # más alta directo en esa columna, pisando la fórmula/valor previo.
        if len(cuotas) > 1:
            _set_cell(ws_deu.cell(srow, cfg['max_cuota_col']), max(cuotas))

    # Partes de cuota. En orden: si un bloque se tocó varias veces en la corrida,
    # el último snapshot ya trae el acumulado (incluidas las cuotas enteras).
    for r in con_partes:
        label = f"{str(r['cliente'])[:38]}{r['lote_str']} " + ' y '.join(r['etiqueta'])
        ws_edit.cell(r['imp_row'], 8).value = label
        for cell in ws_edit[r['imp_row']]:
            cell.fill = YELLOW_FILL
        for b in r['bloques']:
            cfg = sheets_cfg[b['hoja']]
            ws_deu = wb_deu_edit[b['hoja']]
            _set_cell(ws_deu.cell(b['fila'], b['real_col']), b['real'])
            _set_cell(ws_deu.cell(b['fila'], b['cuota_col']), b['cuota'])
            if b['fecha'] is not None:
                _set_cell(ws_deu.cell(b['fila'], b['fecha_col']), b['fecha'])
            if b['max_cuota'] is not None:
                _set_cell(ws_deu.cell(b['fila'], cfg['max_cuota_col']), b['max_cuota'])

    imp_out = io.BytesIO()
    deu_out = io.BytesIO()
    wb_imp.save(imp_out)
    wb_deu_edit.save(deu_out)
    wb_imp.close()
    wb_deu_edit.close()
    return imp_out.getvalue(), deu_out.getvalue()
